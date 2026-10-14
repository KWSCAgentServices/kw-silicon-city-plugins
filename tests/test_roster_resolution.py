"""Synthetic regression cases: no brokerage records or portraits are bundled."""
import importlib.util
import pathlib
import unittest
import tempfile

path = pathlib.Path(__file__).parents[1] / 'skills/kw-silicon-city-plugins/scripts/resolve_roster.py'
spec = importlib.util.spec_from_file_location('resolver', path)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def row(name, category='agent_candidate', **kwargs):
    return dict(name=name, aliases=[name], category=category, **kwargs)


class ResolutionTests(unittest.TestCase):
    def setUp(self):
        self.records = [row('Alex Tran', email='alex@example.test', phone='408-555-0101', license='01234567'),
                        row('Alex Nguyen', email='other@example.test'), row('Sam Lee', 'staff'),
                        row('Taylor Doe', 'former')]

    def test_normalized_name(self):
        self.assertEqual(module.resolve(self.records, ' ALEX  TRAN ')['record']['name'], 'Alex Tran')

    def test_dre_leading_zero(self):
        self.assertEqual(module.resolve(self.records, '01234567')['record']['name'], 'Alex Tran')

    def test_email(self):
        self.assertEqual(module.resolve(self.records, 'ALEX@example.test')['record']['name'], 'Alex Tran')

    def test_phone(self):
        self.assertEqual(module.resolve(self.records, '+1 (408) 555-0101')['record']['name'], 'Alex Tran')

    def test_partial_ambiguous(self):
        self.assertEqual(module.resolve(self.records, 'Alex')['status'], 'needs_identifier')

    def test_partial_unique_is_not_identity(self):
        self.assertEqual(module.resolve(self.records, 'Taylor')['status'], 'needs_identifier')

    def test_staff_role(self):
        self.assertEqual(module.resolve(self.records, 'Sam Lee')['eligibility'], 'not_established_as_current_agent')

    def test_former_role(self):
        self.assertEqual(module.resolve(self.records, 'Taylor Doe')['record']['category'], 'former')

    def test_no_cross_record_fields(self):
        self.assertIsNone(module.resolve(self.records, 'Alex Nguyen')['record'].get('phone'))

    def test_duplicate_identifier(self):
        records = self.records + [row('Duplicate', email='alex@example.test')]
        self.assertEqual(module.resolve(records, 'alex@example.test')['status'], 'ambiguous')

    def test_unknown(self):
        self.assertEqual(module.resolve(self.records, 'Missing Person')['status'], 'not_found')

    def test_blank_license_cannot_match_zero(self):
        self.assertEqual(module.resolve(self.records, '00000000')['status'], 'not_found')

    def test_alias(self):
        records = [dict(name='Alex Tran', aliases=['Alex Tran', 'Alexander Tran'], category='agent_candidate')]
        self.assertEqual(module.resolve(records, 'Alexander Tran')['record']['name'], 'Alex Tran')

    def test_empty_query(self):
        self.assertEqual(module.resolve(self.records, '')['status'], 'not_found')

    def test_no_implicit_license_verification(self):
        self.assertEqual(module.resolve(self.records, 'Alex Tran')['eligibility'], 'verify_license_and_affiliation')

    def test_unknown_private_fields_are_dropped(self):
        records = [row('Alex Tran', home_address='PRIVATE_ADDRESS', spouse='PRIVATE_FAMILY',
                       children=['PRIVATE_CHILD'], private_notes='PRIVATE_NOTES')]
        result = module.resolve(records, 'Alex Tran')
        self.assertNotIn('PRIVATE_', str(result))

    def test_public_projection_is_contact_only(self):
        records = [row('Alex Tran', email='alex@example.test', phone='408-555-0101',
                       license='01234567', expiration='INTERNAL_DATE', role='INTERNAL_ROLE',
                       sheet='INTERNAL_SHEET', row=42, home_address='PRIVATE_ADDRESS')]
        result = module.public_contact_result(module.resolve(records, 'Alex Tran'))
        self.assertEqual(set(result['record']), {'name', 'email', 'phone', 'license'})
        self.assertNotIn('INTERNAL_', str(result))
        self.assertNotIn('PRIVATE_', str(result))

    def test_ambiguous_public_projection_does_not_share_internal_categories(self):
        result = module.public_contact_result(module.resolve(self.records, 'Alex'))
        self.assertEqual(result['status'], 'needs_identifier')
        self.assertTrue(all(set(item) == {'name'} for item in result['candidates']))

    def test_workbook_private_columns_and_unknown_name_headers(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            book = openpyxl.Workbook()
            book.active.title = 'Agent Roster'
            book.active.append(['Agent Name', 'Email', 'Phone', 'DRE', 'Home Address',
                                'Spouse', 'Children', 'Emergency Contact', 'Private Notes'])
            book.active.append(['Alex Tran', 'alex@example.test', '408-555-0101', '01234567',
                                'PRIVATE_ADDRESS', 'PRIVATE_SPOUSE', 'PRIVATE_CHILD',
                                'PRIVATE_EMERGENCY', 'PRIVATE_NOTES'])
            unknown = book.create_sheet('Unknown')
            unknown.append(['Home Address', 'Family Members'])
            unknown.append(['PRIVATE_HOME_AS_NAME', 'PRIVATE_FAMILY'])
            file = pathlib.Path(directory) / 'privacy-fixture.xlsx'
            book.save(file)
            records = module.rows_from_workbook(file)
            self.assertEqual(len(records), 1)
            self.assertNotIn('PRIVATE_', str(records))
            contact = module.public_contact_result(module.resolve(records, 'Alex Tran'))
            self.assertEqual(set(contact['record']), {'name', 'email', 'phone', 'license'})

    def test_sparse_workbook_and_role_sheets(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            book = openpyxl.Workbook()
            book.active.title = 'Updated Roster'
            book.active.append(['Full Name', 'First Name', 'Last Name', 'DRE#'])
            book.active.append(['Alex Tran', 'Alex', 'Tran', '01234567'])
            book.active.append(['Incomplete Row'])
            staff = book.create_sheet('STAFF')
            staff.append(['Full Name', 'Title', 'Email'])
            staff.append(['Sam Lee', 'Agent Services', 'sam@example.test'])
            unknown = book.create_sheet('Miscellaneous')
            unknown.append(['Full Name', 'Lockbox Code', 'Private Notes'])
            unknown.append(['Unverified Person', 'SECRET', 'PRIVATE'])
            file = pathlib.Path(directory) / 'fixture.xlsx'
            book.save(file)
            records = module.rows_from_workbook(file)
            self.assertEqual(module.resolve(records, 'Sam Lee')['record']['category'], 'staff')
            self.assertEqual(module.resolve(records, 'Alex Tran')['record']['license'], '01234567')
            record = module.resolve(records, 'Unverified Person')['record']
            self.assertEqual(record['category'], 'unknown')
            self.assertNotIn('SECRET', str(record))
            self.assertNotIn('PRIVATE', str(record))


if __name__ == '__main__':
    unittest.main()
